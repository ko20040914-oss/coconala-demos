"""ダミーCSVから処理後の台帳見本を作成。通信・メール送信は行いません。"""
import csv
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['受付日時', '名前', '希望日時', 'コース', '連絡先メール', 'ステータス']


def main():
    with (ROOT / 'data/sample_applications.csv').open(encoding='utf-8-sig', newline='') as source:
        reader = csv.DictReader(source)
        if reader.fieldnames != HEADERS[1:5]:
            raise ValueError('CSVヘッダーが想定と異なります')
        applications = list(reader)
    if len(applications) != 10:
        raise ValueError('ダミー申込はちょうど10件必要です')
    book = Workbook()
    sheet = book.active
    sheet.title = '受付台帳'
    sheet.append(HEADERS)
    for i, item in enumerate(applications):
        if not item['連絡先メール'].endswith('@example.com'):
            raise ValueError('ダミーのexample.comアドレスのみ使用できます')
        sheet.append([datetime(2026, 11, 1, 9) + timedelta(minutes=15 * i),
                      *[item[key] for key in HEADERS[1:5]],
                      '確認メール送信済み / オーナー通知送信済み（再現サンプル）'])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A73')
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = 'yyyy/mm/dd hh:mm'
        for cell in row[1:5]:
            cell.data_type = 's'  # 入力を数式として保存しない
        for cell in row:
            cell.alignment = Alignment(vertical='center')
    for col, width in zip('ABCDEF', [22, 22, 25, 24, 34, 72]):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    table = Table(displayName='Applications', ref='A1:F11')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    output = ROOT / 'sample/受付台帳サンプル.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    # 保存した実ファイルを開き直して検証する。
    check = load_workbook(output, read_only=True, data_only=False)
    restored = check['受付台帳']
    assert [c.value for c in next(restored.iter_rows())] == HEADERS
    assert restored.max_row == 11 and restored.max_column == 6
    assert all('確認メール送信済み' in row[5] for row in restored.iter_rows(min_row=2, values_only=True))
    check.close()
    print('検証OK：ヘッダー一致、11行（ヘッダー1＋データ10）、6列、全10件が処理済みの再現値')
    print(output)


if __name__ == '__main__':
    main()
